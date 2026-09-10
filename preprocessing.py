
from pathlib import Path
from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_core.documents import Document
import torch

import docx
from pptx import Presentation

import pytesseract
from PIL import Image

import whisper
from transformers import pipeline
import os


os.environ["PATH"] += os.pathsep + r"C:\Users\Adam\AppData\Local\Microsoft\WinGet\Packages\Gyan.FFmpeg_Microsoft.Winget.Source_8wekyb3d8bbwe\ffmpeg-8.0.1-full_build\bin"



# -------------------------
# DOCUMENT PROCESSOR
# -------------------------

# --- Pipeline ---
def document_pipeline(path):
    result = []
    
    print("Step 1: Loading document...")
    loader = read_document_by_type(path)
    print("Step 2: Generating chunks...")
    chunks = generate_chunks(loader)
    print("Step 3: Generating embeddings...")
    embeddings = generate_embeddings(chunks)
    print("Step 4: Combining chunks and embeddings...")
    for i in range(len(chunks)):
        chunk_doc = chunks[i]
        
        if i < len(embeddings) and isinstance(embeddings[i], list) and len(embeddings[i]) > 0:
            emb = embeddings[i]
        else:
        
            emb = [0.0] * 384  
            
        result.append({
            "chunk": chunk_doc.page_content,
            "embedding": emb,
            "metadata": chunk_doc.metadata
        })
    
    print(f"Document processing complete! Created {len(result)} processed chunks")
    return result

# --- Read by Type --- 
def read_document_by_type(path):
    p = Path(path)
    file_type = p.suffix.lower()
    action_map = {
        ".pdf": read_pdf,
        ".pptx": read_pptx,
        ".docx": read_docx,
        ".md": read_txt_md,
        ".txt": read_txt_md,
        # ".wav": read_speech,
        # ".png": read_images,
        # ".jpg": read_images,   
    }
    
    if file_type not in action_map:
        raise ValueError(f"Unsupported file type: {file_type}")
    return action_map[file_type](path) 

# --- Chunker ---
def generate_chunks(lang_docs):
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1200,
        chunk_overlap=200,
        separators=["\n\n", "\n", ".", "!", "?"],
        length_function=len
    )
    return splitter.split_documents(lang_docs)

# --- Embedder ---
def generate_embeddings(chunks, model_name = "BAAI/bge-small-en-v1.5"):
    device = "cuda" if torch.cuda.is_available() else "cpu"
    embedding_model = HuggingFaceEmbeddings(
        model_name=model_name,
        model_kwargs={"device": device},
    )
    
    texts = [chunk.page_content for chunk in chunks]
    embeddings = []

    batch_size = 64  
    for i in range(0, len(texts), batch_size):
        batch = texts[i:i+batch_size]
        emb_batch = embedding_model.embed_documents(batch)
        embeddings.extend(emb_batch)
    
    return embeddings

def generate_query_embedding(query: str, model_name="BAAI/bge-small-en-v1.5"):

    device = "cuda" if torch.cuda.is_available() else "cpu"

    embedding_model = HuggingFaceEmbeddings(
        model_name=model_name,
        model_kwargs={"device": device},
    )
    
    query_embedding = embedding_model.embed_query(query)
    
    return query_embedding

# --- Text Cleaner ---
def text_formatter(text: str) -> str:
    return text.replace("\n", " ").strip()

# --- PDF Loader ---
def read_pdf(path):
    loader = PyPDFLoader(path)
    docs = loader.load()
    allowed = ["total_pages", "page", "source", "title"]
    for d in docs:
        d.metadata = {k: v for k, v in d.metadata.items() if k in allowed}
    return docs

# --- DOCX loader ---
def read_docx(path):
    document = docx.Document(path)
    docs = []
    for i, p in enumerate(document.paragraphs):
        if p.text.strip():
            docs.append(Document(
                page_content=text_formatter(p.text),
                metadata={"page": i+1, "source": path}
            ))
    return docs

# --- PPTX loader ---
def read_pptx(path):
    pres = Presentation(path)
    docs = []
    for i, slide in enumerate(pres.slides):
        text = " ".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))
        if text.strip():
            docs.append(Document(
                page_content=text_formatter(text),
                metadata={"page": i+1, "source": path}
            ))
    return docs

# --- TXT / Markdown loader ---
def read_txt_md(path):
    docs = []
    with open(path, "r", encoding="utf-8") as f:
        text = f.read()
        docs.append(Document(
            page_content=text_formatter(text),
            metadata={"page": 1, "source": path}
        ))
    return docs

# --- OCR ---
def read_image(image_path: str) -> str:
    pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    
    try:
        image = Image.open(image_path)
        text = pytesseract.image_to_string(image)
    except pytesseract.TesseractNotFoundError:
        raise RuntimeError(
            f"Tesseract is not installed or not found at the specified path. "
            f"Please install Tesseract-OCR and set the correct path. "
            f"Current path: {pytesseract.pytesseract.tesseract_cmd}"
        )
    except FileNotFoundError:
        raise RuntimeError(f"Image file not found: {image_path}")
    except Exception as e:
        raise RuntimeError(f"An error occurred during OCR processing: {e}")

    text = text_formatter(text)
    
    # Check if OCR returned any text
    if not text.strip():
        print(f"Warning: No text detected in image {image_path}")
        return ""
    
    return text

# --- Transcribe Audio ---
_model_cache = {}

def read_audio(audio_path: str, model_size: str = "base") -> str:
    if model_size not in _model_cache:
        _model_cache[model_size] = whisper.load_model(model_size)

    model = _model_cache[model_size]
    result = model.transcribe(audio_path)
    return result["text"].strip()

